# -*- coding: utf-8 -*-
"""자료 폴더의 재료를 한 파일로 모은다.

쓰는 법:  python gather.py 자료 작업/재료.md
하는 일
  1. 자료 폴더에서 회사 양식을 고른다. 이름에 양식/템플릿이 있는 것, 그다음 가로 쪽이 대부분인 PDF,
     없으면 파워포인트 파일(pptx, potx). 파워포인트 파일이면 파워포인트로 PDF를 만든다. 고른 것을 첫 줄에 알린다.
  2. 나머지 파일을 읽어 작업/재료.md 한 파일에 적는다. 숫자 검사는 이 파일과 대조한다.
     글(md, txt), 표(csv, xlsx), 워드(docx), 다른 PDF, 파워포인트(pptx)의 글.
     한글(hwp)은 읽지 못한다. 사람에게 PDF로 저장해 넣어 달라고 보고에 적는다.
  3. 그림 파일(png, jpg, svg)은 목록만 적는다.
엑셀 숫자는 엑셀 화면에 보이는 모양(천 단위 쉼표, 백분율, 소수 자리)으로 적는다.
"""
import sys, io, os, re, csv, zipfile, datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
src, out = sys.argv[1], sys.argv[2]
IMG = {".png", ".jpg", ".jpeg", ".svg", ".gif", ".bmp", ".webp", ".emf"}


def read_text(p):
    raw = open(p, "rb").read()
    for enc in ("utf-8-sig", "cp949", "utf-16"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", "replace")


def md_table(rows):
    rows = [[("" if c is None else str(c)).replace("|", "/").replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    keep = [j for j in range(w) if any(j < len(r) and r[j] for r in rows)]
    rows = [[r[j] if j < len(r) else "" for j in keep] for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(keep)]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def fmt_cell(c):
    v = c.value
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    f = c.number_format or "General"
    dec = 0
    m = re.search(r"0\.(0+)", f)
    if m:
        dec = len(m.group(1))
    if "%" in f:
        return f"{v * 100:.{dec}f}%"
    if f == "General":
        if float(v).is_integer():
            return str(int(v))
        return f"{v:.4f}".rstrip("0").rstrip(".")
    s = f"{v:,.{dec}f}" if "," in f else f"{v:.{dec}f}"
    return s


def read_xlsx(p):
    try:
        import openpyxl
    except ImportError:
        return "(엑셀을 읽으려면 먼저 설치: pip install openpyxl)"
    wb = openpyxl.load_workbook(p, data_only=True)
    parts = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows = [[fmt_cell(c) for c in r] for r in ws.iter_rows()]
        t = md_table(rows)
        if t:
            merged = ", ".join(str(r) for r in list(ws.merged_cells.ranges)[:10])
            parts.append(f"### 시트: {ws.title}\n" + (f"(병합된 칸: {merged})\n" if merged else "") + "\n" + t)
    return "\n\n".join(parts)


def read_csv(p):
    return md_table(list(csv.reader(io.StringIO(read_text(p)))))


def read_docx(p):
    try:
        import docx
    except ImportError:
        return "(워드를 읽으려면 먼저 설치: pip install python-docx)"
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for t in d.tables:
        parts.append(md_table([[c.text for c in r.cells] for r in t.rows]))
    return "\n\n".join(parts)


def read_pdf(p):
    import fitz
    d = fitz.open(p)
    return "\n\n".join(f"(p.{i + 1})\n{pg.get_text().strip()}" for i, pg in enumerate(d))


def read_pptx(p):
    z = zipfile.ZipFile(p)
    names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=lambda n: int(re.findall(r"\d+", n)[-1]))
    parts = []
    for i, n in enumerate(names):
        x = z.read(n).decode("utf-8", "replace")
        paras = ["".join(re.findall(r"<a:t>([^<]*)</a:t>", pp)) for pp in re.findall(r"<a:p>.*?</a:p>", x, re.S)]
        parts.append(f"(슬라이드 {i + 1})\n" + "\n".join(t for t in paras if t.strip()))
    return "\n\n".join(parts)


files = []
for root, _, names in os.walk(src):
    for n in sorted(names):
        if not n.startswith(("~$", ".")):
            files.append(os.path.join(root, n))

# 회사 양식 고르기: 이름에 양식/템플릿이 있는 것, 그다음 가로 쪽이 대부분인 PDF, 없으면 파워포인트 파일
NAMED = ("양식", "템플릿", "template")
template, best = None, -1
import fitz
for p in files:
    low = p.lower()
    named = any(k in os.path.basename(low) for k in NAMED)
    if low.endswith(".pdf"):
        d = fitz.open(p)
        land = sum(1 for pg in d if pg.rect.width > pg.rect.height)
        if land * 2 < len(d) and not named:
            continue  # 세로 문서는 재료로 본다
        score = 2 * 10 ** 6 * named + 10 ** 5 + len(d)
    elif low.endswith((".pptx", ".potx")):
        n = sum(1 for x in zipfile.ZipFile(p).namelist() if re.match(r"ppt/slides/slide\d+\.xml$", x))
        score = 2 * 10 ** 6 * named + n
    else:
        continue
    if score > best:
        template, best = p, score
template_pdf = template
if template and template.lower().endswith((".pptx", ".potx")):
    # 파워포인트 파일이면 파워포인트로 PDF를 만들어 양식으로 쓴다
    work = os.path.dirname(os.path.abspath(out))
    src_pptx = template
    if template.lower().endswith(".potx"):
        import shutil
        src_pptx = os.path.join(work, "회사양식_사본.pptx")
        shutil.copy(template, src_pptx)
    import subprocess
    subprocess.run([sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "render.py"), src_pptx, work, "회사양식"], capture_output=True)
    template_pdf = os.path.join(work, "회사양식.pdf")
    import glob
    for f in glob.glob(os.path.join(work, "회사양식_s*.png")) + [os.path.join(work, "회사양식_sheet.jpg")]:
        if os.path.exists(f):
            os.remove(f)
    if not os.path.exists(template_pdf):
        template_pdf = None

out_parts, images, skipped = [], [], []
for p in files:
    ext = os.path.splitext(p)[1].lower()
    rel = os.path.relpath(p, src).replace("\\", "/")
    if p == template:
        continue
    if ext in IMG:
        images.append(rel)
        continue
    try:
        if ext in (".md", ".txt"):
            body = read_text(p)
        elif ext == ".csv":
            body = read_csv(p)
        elif ext in (".xlsx", ".xlsm"):
            body = read_xlsx(p)
        elif ext == ".docx":
            body = read_docx(p)
        elif ext == ".pdf":
            body = read_pdf(p)
        elif ext == ".pptx":
            body = read_pptx(p)
        else:
            skipped.append(f"{rel} (읽지 못하는 형식. PDF나 txt로 저장해 넣어 달라고 보고에 적는다)")
            continue
    except Exception as e:
        skipped.append(f"{rel} (읽다가 실패: {e})")
        continue
    out_parts.append(f"## 파일: {rel}\n\n{body.strip()}\n")

os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
head = ["# 재료 모음 (gather.py가 자료 폴더에서 모은 것)", ""]
if images:
    head.append("그림 파일: " + ", ".join(images))
if skipped:
    head.append("읽지 못한 파일: " + "; ".join(skipped))
open(out, "w", encoding="utf-8").write("\n".join(head) + "\n\n" + "\n".join(out_parts))

rel = lambda p: os.path.relpath(p, ".").replace(chr(92), "/")
if template_pdf and template_pdf != template:
    print(f"회사 양식: {rel(template_pdf)} ({rel(template)}를 파워포인트로 PDF로 바꾼 것. 다음 단계에는 이 PDF를 쓴다)")
elif template_pdf:
    print(f"회사 양식: {rel(template_pdf)}")
else:
    print("회사 양식: (찾지 못함. 자료 폴더에 회사 발표 자료(PPT 또는 PDF)를 넣어 달라고 보고한다)")
print(f"재료 {len(out_parts)}개를 {out}에 모았다. 그림 {len(images)}개" + (f", 읽지 못한 파일 {len(skipped)}개" if skipped else ""))
for s in skipped:
    print("  읽지 못함:", s)
